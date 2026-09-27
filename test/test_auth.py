from openpyxl import load_workbook
from playwright.sync_api import Page

from page import auth_page


def test_valid_login(page: Page):
    wb = load_workbook("testdata.xlsx")
    sheet = wb["valid_login"]

    for i in range(2, sheet.max_row + 1):
        auth_page.valid_login(page, sheet[f"A{i}"].value, sheet[f"B{i}"].value)
        auth_page.logout(page)


def test_invalid_login(page: Page):
    wb = load_workbook("testdata.xlsx")
    sheet = wb["invalid_login"]

    for i in range(2, sheet.max_row + 1):
        auth_page.invalid_login(page, sheet[f"A{i}"].value, sheet[f"B{i}"].value)
